from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "Hexa avocado Informe Final.docx"
OUT = ROOT / "assets" / "images"
DATA = ROOT / "content.json"
MAP = ROOT / "image-map.md"

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def slug_ext(name: str) -> str:
    suffix = Path(name).suffix.lower()
    return suffix if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"} else ".png"


def paragraph_text(paragraph) -> str:
    return re.sub(r"\s+", " ", paragraph.text).strip()


def iter_block_items(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            for p in document.paragraphs:
                if p._element is child:
                    yield ("p", p)
                    break
        elif child.tag.endswith("}tbl"):
            for t in document.tables:
                if t._element is child:
                    yield ("table", t)
                    break


def table_text(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        rows.append([re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells])
    return rows


def rels_for_docx(zf: zipfile.ZipFile) -> dict[str, str]:
    rel_root = ET.fromstring(zf.read("word/_rels/document.xml.rels"))
    return {node.attrib["Id"]: node.attrib["Target"] for node in rel_root}


def image_ids_from_paragraph(paragraph) -> list[str]:
    ids = []
    for blip in paragraph._element.findall(".//a:blip", NS):
        rid = blip.attrib.get(f"{{{NS['r']}}}embed")
        if rid:
            ids.append(rid)
    return ids


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)

    with zipfile.ZipFile(SOURCE) as zf:
        rels = rels_for_docx(zf)
        blocks = []
        images = []
        pending_caption = None
        image_count = 0
        previous_text = ""

        for kind, item in iter_block_items(doc):
            if kind == "table":
                rows = table_text(item)
                if any(any(cell for cell in row) for row in rows):
                    blocks.append({"type": "table", "rows": rows})
                continue

            text = paragraph_text(item)
            image_ids = image_ids_from_paragraph(item)

            if text:
                blocks.append({"type": "paragraph", "text": text})
                if re.match(r"^(Imagen|Figura|Layout)\s*\d+", text, re.I):
                    pending_caption = text
                else:
                    previous_text = text

            for rid in image_ids:
                target = rels[rid]
                src = "word/" + target if not target.startswith("word/") else target
                image_count += 1
                name = f"image-{image_count:02d}{slug_ext(src)}"
                dest = OUT / name
                with zf.open(src) as source, dest.open("wb") as output:
                    shutil.copyfileobj(source, output)

                image = {
                    "number": image_count,
                    "file": f"assets/images/{name}",
                    "caption": pending_caption or f"Imagen {image_count}",
                    "after": previous_text,
                }
                images.append(image)
                blocks.append({"type": "image", **image})
                pending_caption = None

    DATA.write_text(json.dumps({"blocks": blocks, "images": images}, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = ["# Mapa de imágenes", ""]
    for image in images:
        lines.extend(
            [
                f"## Imagen {image['number']}",
                f"- Archivo: `{image['file']}`",
                f"- Título en Word: {image['caption']}",
                f"- Va después del texto: {image['after'] or 'No se detectó texto previo.'}",
                "",
            ]
        )
    MAP.write_text("\n".join(lines), encoding="utf-8")
    print(f"Extraídas {len(images)} imágenes.")


if __name__ == "__main__":
    main()
